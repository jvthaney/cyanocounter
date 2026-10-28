#!/usr/bin/env python3
"""
build_dashboard.py

Writes a self-contained food web dashboard: one HTML file that carries the trait
lookup and the feeding rules inside it, so a user can load a count, edit it, add or
remove taxa, change the rule parameters, and watch the web redraw. No server, no
Python, nothing to install. Open the file, or publish it beside the counters.

    conda activate video
    python build_dashboard.py \\
        --zooplankton "$ZK/data/cfbkey_workbook.xlsx" \\
        --phytoplankton "$ZK/data/phyto_traits.xlsx" \\
        --rules "$ZK/data/foodweb_rules.csv" \\
        --out "$ZK/docs/foodweb.html"

WHAT THE PAGE DOES
  * Loads a CSV exported by FoodWebCounter (taxon, group, counting_unit, count).
  * Lets the user edit counts, change a counting unit, and add or remove taxa, with
    name completion from the trait tables.
  * Applies the same rules as build_food_web.py, in the browser, and redraws.
  * Reports degree centrality, betweenness centrality, connectance and link count.
  * Node colour is the group, node size is abundance, so a taxon that is both
    abundant and well connected stands out.

WHAT TRAVELS WITH IT
  Only the fields the rules need: for an animal its body length, feeding mode and
  mesh class; for an alga its largest dimension and whether it is a colony or a
  filament. No taxonomy, no references, no images. The file is a few hundred
  kilobytes.

HONEST LIMITS, SHOWN ON THE PAGE ITSELF
  A taxon with no body length cannot be linked. A link means "physically possible
  on these rules", not "observed". The rule parameters that have no published
  source are marked as such in the panel where they are edited.
"""

import argparse
import csv
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Needs openpyxl:  conda install -c conda-forge openpyxl")

MESH_CLASS = {"fine": ("Diaphanosoma", "Ceriodaphnia", "Chydorus"),
              "coarse": ("Holopedium", "Sida", "Latona", "Latonopsis")}
MESH_SPECIES = {"Daphnia cucullata": "fine", "Daphnia magna": "fine"}

DEFAULT_RULES = {
    "burns_slope": 22, "burns_intercept": 4.87,
    "mesh_fine_um": 0.5, "mesh_medium_um": 1.0, "mesh_coarse_um": 3.0, "mesh_default_um": 1.0,
    "interference_length_um": 100, "raptorial_particle_multiplier": 2.0,
    "prey_ratio_min": 0.05, "prey_ratio_max": 0.50, "rotifer_max_particle_um": 20,
}
RULE_NOTES = {
    "burns_slope": "Burns 1968, published",
    "burns_intercept": "Burns 1968, published",
    "mesh_fine_um": "interpolated from Brendelberger's 0.2-4.7 um range, not published",
    "mesh_medium_um": "interpolated, not published",
    "mesh_coarse_um": "interpolated, not published",
    "mesh_default_um": "placeholder for rotifers and copepods, which do not sieve",
    "interference_length_um": "OUR PLACEHOLDER: no verified published threshold",
    "raptorial_particle_multiplier": "NO SOURCE: a plausible guess",
    "prey_ratio_min": "predator:prey size window, rule of thumb",
    "prey_ratio_max": "predator:prey size window, rule of thumb",
    "rotifer_max_particle_um": "reasonable for microphagous rotifers, wrong for Asplanchna",
}


def rows_of(path, sheet):
    wb = load_workbook(Path(path).expanduser(), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    h = [str(c).strip() if c is not None else "" for c in rows[0]]
    return [dict(zip(h, r)) for r in rows[1:] if any(x is not None for x in r)]


def get(d, *names):
    low = {str(k).strip().lower(): v for k, v in d.items() if k}
    for n in names:
        v = low.get(n.lower())
        if v not in (None, ""):
            return v
    return ""


def fnum(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"\d+(?:\.\d+)?", str(v).replace(",", "."))
    return float(m.group(0)) if m else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--zooplankton", required=True)
    ap.add_argument("--phytoplankton", required=True)
    ap.add_argument("--rules", default="")
    ap.add_argument("--aliases", default="",
                    help="the name-difference CSV (rimet_name,phycokey_name,note), so a "
                         "current name such as Dolichospermum finds Anabaena in the table")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    rules = dict(DEFAULT_RULES)
    if a.rules and Path(a.rules).expanduser().exists():
        for line in Path(a.rules).expanduser().read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or line.lower().startswith("parameter,"):
                continue
            bits = next(csv.reader([line]))
            if len(bits) >= 2 and fnum(bits[1]) is not None and bits[0].strip() in rules:
                rules[bits[0].strip()] = fnum(bits[1])

    # ---- animals ----
    animals = []
    for r in rows_of(a.zooplankton, "feeding_groups"):
        name = str(get(r, "display_name", "species_id")).strip()
        if not name:
            continue
        lo = fnum(get(r, "body_length_min_mm"))
        hi = fnum(get(r, "body_length_max_mm")) or lo
        if lo is None:
            # fall back to the length mined from the key's own text, as
            # build_food_web.py does: "0.8-2.2" or a single "0.5"
            draft = str(get(r, "body_length_mm_draft"))
            nums = re.findall(r"\d+(?:\.\d+)?", draft.replace(",", "."))
            if nums:
                lo = float(nums[0])
                hi = float(nums[1]) if len(nums) > 1 else lo
        genus = str(get(r, "genus")).strip()
        mesh = "medium"
        for cls, gs in MESH_CLASS.items():
            if genus in gs:
                mesh = cls
        mesh = MESH_SPECIES.get(name, mesh)
        animals.append({
            "n": name, "g": "zoop",
            "lo": lo, "hi": hi, "mesh": mesh,
            "m": str(get(r, "feeding_mode", "feeding_mode_draft")).strip().lower(),
            "t": str(get(r, "trophic_group", "trophic_group_draft")).strip().lower(),
            "r": 1 if str(get(r, "phylum", "Phylum")).lower().startswith("rotifer") else 0,
        })

    # ---- algae, one entry per counting unit ----
    algae = []
    for r in rows_of(a.phytoplankton, "phyto_traits"):
        role = str(get(r, "role")).strip().lower()
        if role not in ("phytoplankton", "uncertain", ""):
            continue      # macroalgae, protozoa and the anomalous items are not food here
        size = fnum(get(r, "max_dimension_um"))
        if size is None:
            continue
        unit = str(get(r, "counting_unit")).strip()
        morph = str(get(r, "morphology_phycokey")).strip().lower()
        algae.append({
            "n": str(get(r, "taxon")).strip(), "g": "phyto", "u": unit or "cell",
            "s": size, "bv": fnum(get(r, "biovolume_um3")),
            "f": 1 if (unit == "filament_100um" or "filament" in morph) else 0,
            "c": 1 if unit == "colony" else 0,
            "gen": str(get(r, "genus")).strip(),
            "gs": 1 if str(get(r, "row_type")).strip() == "genus_summary" else 0,
            "ns": fnum(get(r, "n_species_in_genus")),
        })

    aliases = {}
    if a.aliases and Path(a.aliases).expanduser().exists():
        for line in Path(a.aliases).expanduser().read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or line.lower().startswith("rimet_name"):
                continue
            bits = next(csv.reader([line]))
            if len(bits) >= 2 and bits[0].strip() and bits[1].strip():
                aliases[bits[0].strip().lower()] = bits[1].strip()
                aliases.setdefault(bits[1].strip().lower(), bits[0].strip())
    if aliases:
        print(f"  name aliases loaded: {len(aliases)}")

    payload = {"animals": animals, "algae": algae, "rules": rules, "notes": RULE_NOTES,
               "aliases": aliases}
    html = PAGE.replace("/*__DATA__*/", json.dumps(payload, separators=(",", ":")))
    outp = Path(a.out).expanduser()
    outp.parent.mkdir(parents=True, exist_ok=True)
    outp.write_text(html, encoding="utf-8")
    size_kb = outp.stat().st_size / 1024
    with_len = sum(1 for x in animals if x["hi"])
    print(f"Wrote {outp}  ({size_kb:.0f} KB)")
    print(f"  animals in the lookup: {len(animals)}  ({with_len} with a body length)")
    print(f"  algal counting units:  {len(algae)}")
    print("\nOpen it in a browser, or publish it beside the counters.")


PAGE = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Food web dashboard</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css?family=Cabin|News Cycle">
<script src="https://cdn.jsdelivr.net/npm/d3@7.8.5/dist/d3.min.js"></script>
<style>
 :root { --phyto:#2e8b57; --zoop:#F8BD7F; --other:#727171; --ink:#222; --line:#ccc; }
 body { font-family:'Cabin',sans-serif; color:var(--ink); margin:0; background:#fafaf8; }
 header { background:#333; color:#fff; padding:14px 20px; }
 header h1 { font-family:'News Cycle',sans-serif; font-size:26px; margin:0; color:lightgreen; }
 header p { margin:4px 0 0; font-size:14px; color:#ddd; }
 header nav { margin-top:8px; }
 header nav a { color:#fff; font-size:15px; text-decoration:none; padding:3px 10px;
                border:1px solid rgba(255,255,255,0.4); border-radius:5px;
                display:inline-block; margin:2px 4px 2px 0; }
 header nav a:hover { background:rgba(255,255,255,0.15); }
 header nav a[aria-current="page"] { background:#fff; color:#222; border-color:#fff; }
 .wrap { display:grid; grid-template-columns:minmax(360px,460px) minmax(0,1fr); gap:20px;
         padding:18px 20px; align-items:start; }
 @media (max-width:1000px){ .wrap { grid-template-columns:1fr; } }
 .panel { background:#fff; border:1px solid var(--line); border-radius:8px; padding:14px; }
 .panel h2 { font-size:16px; text-transform:uppercase; letter-spacing:.06em; margin:0 0 10px;
             background:#333; color:#fff; padding:6px 10px; border-radius:4px; }
 table { border-collapse:collapse; width:100%; font-size:14px; }
 th,td { border:1px solid #ddd; padding:4px 6px; text-align:left; }
 th { background:#f0f0ee; }
 input,select,button { font-family:'Cabin',sans-serif; font-size:14px; }
 input[type=number] { width:76px; }
 .taxon-in { width:100%; min-width:120px; }
 button { cursor:pointer; border:1px solid #444; background:#fff; border-radius:5px; padding:4px 10px; }
 button.primary { background:#1f6f9c; color:#fff; border-color:#1f6f9c; }
 button:focus-visible,input:focus-visible,select:focus-visible { outline:3px solid #ffbf47; outline-offset:1px; }
 .rm { color:#a00; border-color:#a00; padding:2px 7px; }
 #graph { width:100%; height:620px; background:#fff; border:1px solid var(--line); border-radius:8px; }
 .legend span { display:inline-block; margin-right:14px; font-size:14px; }
 .sw { display:inline-block; width:12px; height:12px; border-radius:50%; margin-right:5px;
        vertical-align:middle; box-shadow:0 0 0 1px rgba(0,0,0,0.25); }
 .warn { background:#fff8e1; border-left:4px solid #EE7733; padding:8px 12px; font-size:14px; margin:8px 0; }
 .notelist { max-height:170px; overflow-y:auto; margin:4px 0; }
 .notelist div { padding:1px 0; }
 .muted { color:#666; font-size:13px; }
 .from-counter { background:#eaf4ec; border-left:4px solid #2e8b57; padding:8px 12px;
                 font-size:14px; margin:8px 0; }
 details summary { cursor:pointer; font-weight:bold; margin:6px 0; }
 .stat { font-size:14px; }
 .stat b { font-size:20px; }
 node text { pointer-events:none; }
</style>
</head>
<body>
<header>
  <h1>Food web dashboard</h1>
  <p>Load a count from FoodWebCounter, or build one by hand. Edit anything and the web redraws.</p>
  <nav aria-label="Counters and tools">
    <a href="index.html">CyanoCounter</a>
    <a href="Plankton.html">PlanktonCounter</a>
    <a href="Taxon.html">TaxonCounter</a>
    <a href="FoodWeb.html">FoodWebCounter</a>
    <a href="dashboard.html" aria-current="page">Food web dashboard</a>
  </nav>
</header>

<div class="wrap">
 <div>
  <div class="panel">
    <h2>1. The count</h2>
    <p class="muted">Load a CSV exported by FoodWebCounter, or add taxa by hand. Start typing a
       name for suggestions from the trait tables.</p>
    <input type="file" id="file" accept=".csv">
    <button onclick="addRow()">Add taxon</button>
    <button onclick="loadExample()">Load an example</button>
    <button class="rm" onclick="clearAll()">Clear</button>
    <div id="handoff-note"></div>
    <div id="count-warn"></div>
    <table id="count-table">
      <thead><tr><th>Taxon</th><th>Group</th><th>Unit</th><th>Count</th><th></th></tr></thead>
      <tbody id="count-body"></tbody>
    </table>
    <datalist id="taxon-list"></datalist>
  </div>

  <div class="panel" style="margin-top:16px;">
    <h2>2. The rules</h2>
    <p class="muted">Change a number and the web redraws. The note says how much to trust it.</p>
    <table><tbody id="rule-body"></tbody></table>
    <button onclick="resetRules()" style="margin-top:8px;">Reset to defaults</button>
  </div>

  <div class="panel" style="margin-top:16px;">
    <h2>3. Centrality</h2>
    <p class="muted">Degree is how many links a taxon has. Betweenness is how often it sits on the
       shortest path between two others: a high value means energy routes through it.</p>
    <table><thead><tr><th>Taxon</th><th>Degree</th><th>Betweenness</th></tr></thead>
      <tbody id="cent-body"></tbody></table>
  </div>
 </div>

 <div>
  <div class="panel">
    <h2>The web</h2>
    <div class="legend">
      <span><i class="sw" style="background:var(--phyto)"></i>Phytoplankton</span>
      <span><i class="sw" style="background:var(--zoop)"></i>Zooplankton</span>
      <span><i class="sw" style="background:var(--other)"></i>Other</span>
      <span class="muted">node size = abundance</span>
    </div>
    <div id="summary" style="margin:8px 0;"></div>
    <svg id="graph" role="img" aria-label="Food web diagram"></svg>
    <div id="detail" class="muted" style="margin-top:8px;">Select a node to see its links.</div>
    <details>
      <summary>What a link means, and what it does not</summary>
      <p class="muted">A link means the rules make feeding <b>physically possible</b>: the food falls
      inside the consumer's size window and is not a filament or a large colony. It does not mean the
      interaction has been observed. Size rules over-predict, so treat the web as a set of candidate
      links to check, not a finding. A taxon with no body length in the trait table cannot be linked
      at all and is drawn without edges.</p>
    </details>
  </div>
 </div>
</div>

<script>
const DATA = /*__DATA__*/;
let rules = Object.assign({}, DATA.rules);
let rows  = [];          // { taxon, group, unit, count }
let nextRow = 1;

const byName = {};
DATA.animals.forEach(a => { byName[a.n.toLowerCase()] = a; });
const algaeByName = {}, algaeByGenus = {};
DATA.algae.forEach(p => {
  const k = p.n.toLowerCase();
  (algaeByName[k] = algaeByName[k] || []).push(p);
  if (p.gen) { const g = p.gen.toLowerCase();
               (algaeByGenus[g] = algaeByGenus[g] || []).push(p); }
});
const animalsByGenus = {};
DATA.animals.forEach(a => {
  const g = (a.n || '').trim().split(/\s+/)[0].toLowerCase();
  if (g) (animalsByGenus[g] = animalsByGenus[g] || []).push(a);
});

// A counter name often carries the unit in it ("Synura Colony") and is often just
// the genus ("Dinobryon"). Both are handled rather than reported as unknown.
const UNIT_WORDS = { colony:'colony', colonies:'colony', filament:'filament_100um',
                     filaments:'filament_100um', cell:'cell', cells:'cell',
                     'single cell':'cell', trichome:'filament_100um' };

function splitUnitFromName(raw) {
  let name = (raw || '').trim(), unit = null;
  const m = name.match(/[\s,(]+(single cell|colonies|colony|filaments?|cells?|trichome)s?\)?$/i);
  if (m) {
    unit = UNIT_WORDS[m[1].toLowerCase()] || null;
    name = name.slice(0, m.index).trim();
  }
  return { name, unit };
}

// name completion
(function fillList(){
  const dl = document.getElementById('taxon-list');
  const names = [...new Set([...DATA.animals.map(a=>a.n), ...DATA.algae.map(p=>p.n)])].sort();
  dl.innerHTML = names.slice(0, 4000).map(n=>`<option value="${n.replace(/"/g,'&quot;')}">`).join('');
})();

// ─── the count table ──────────────────────────────────────────────────────────
function addRow(taxon, group, unit, count) {
  const id = nextRow++;
  rows.push({ id, taxon: taxon||'', group: group||'phyto', unit: unit||'cell', count: count||0 });
  const tr = document.createElement('tr');
  tr.id = 'r'+id;
  tr.innerHTML = `
    <td><input class="taxon-in" list="taxon-list" value="${(taxon||'').replace(/"/g,'&quot;')}"
         oninput="setField(${id},'taxon',this.value)" aria-label="Taxon name"></td>
    <td><select onchange="setField(${id},'group',this.value)" aria-label="Group">
          <option value="phyto">Phyto</option><option value="zoop">Zoop</option>
          <option value="other">Other</option></select></td>
    <td><select onchange="setField(${id},'unit',this.value)" aria-label="Counting unit" id="u${id}"></select></td>
    <td><input type="number" min="0" value="${count||0}"
         oninput="setField(${id},'count',this.value)" aria-label="Count"></td>
    <td><button class="rm" onclick="removeRow(${id})" aria-label="Remove">&#10005;</button></td>`;
  document.getElementById('count-body').appendChild(tr);
  tr.querySelector('select').value = group || 'phyto';
  renderUnits(id);
  redraw();
}

function renderUnits(id) {
  const r = rows.find(r=>r.id===id); if (!r) return;
  const sel = document.getElementById('u'+id);
  const opts = r.group === 'phyto' ? ['cell','colony','filament_100um'] : ['individual'];
  sel.innerHTML = opts.map(o=>`<option value="${o}">${o==='filament_100um'?'filament':o}</option>`).join('');
  if (!opts.includes(r.unit)) r.unit = opts[0];
  sel.value = r.unit;
}

function setField(id, field, value) {
  const r = rows.find(r=>r.id===id); if (!r) return;
  r[field] = field === 'count' ? Math.max(0, parseFloat(value)||0) : value;
  if (field === 'group') renderUnits(id);
  redraw();
}

function removeRow(id) {
  rows = rows.filter(r=>r.id!==id);
  document.getElementById('r'+id)?.remove();
  redraw();
}

function clearAll() {
  rows = []; document.getElementById('count-body').innerHTML = ''; redraw();
}

function loadExample() {
  clearAll();
  [['Microcystis aeruginosa','phyto','colony',40],
   ['Microcystis aeruginosa (single cell)','phyto','cell',120],
   ['Cyclotella','phyto','cell',200],
   ['Asterionella formosa','phyto','cell',150],
   ['Daphnia pulex','zoop','individual',18],
   ['Bosmina longirostris','zoop','individual',44],
   ['Keratella cochlearis','zoop','individual',60],
   ['Mesocyclops edax','zoop','individual',7],
   ['Leptodora kindtii','zoop','individual',2]].forEach(r=>addRow(...r));
}

// ─── the rules panel ──────────────────────────────────────────────────────────
function renderRules() {
  document.getElementById('rule-body').innerHTML = Object.keys(rules).map(k=>`
    <tr><td><label for="rule-${k}">${k.replace(/_/g,' ')}</label></td>
        <td><input type="number" step="any" id="rule-${k}" value="${rules[k]}"
             oninput="setRule('${k}',this.value)"></td>
        <td class="muted">${DATA.notes[k]||''}</td></tr>`).join('');
}
function setRule(k, v) { const n = parseFloat(v); if (!isNaN(n)) { rules[k] = n; redraw(); } }
function resetRules() { rules = Object.assign({}, DATA.rules); renderRules(); redraw(); }

// ─── the rules themselves, matching build_food_web.py ─────────────────────────
function lookupAnimal(name) { return byName[(name||'').trim().toLowerCase()] || null; }

function pickUnit(list, unit) {
  return list.find(p => p.u === unit) || list.find(p => p.u === 'cell') || list[0];
}

// Returns { trait, how } so the page can say HOW a name was matched.
function lookupAlga(rawName, unit) {
  const exact = algaeByName[(rawName||'').trim().toLowerCase()];
  if (exact) return { trait: pickUnit(exact, unit), how: 'name' };

  const { name, unit: unitFromName } = splitUnitFromName(rawName);
  const u = unitFromName || unit;
  const stripped = algaeByName[name.toLowerCase()];
  if (stripped) return { trait: pickUnit(stripped, u), how: 'name without the unit word' };

  // the genus: prefer the genus median row, which carries the spread across species
  // a different name for the same alga (Dolichospermum / Anabaena)
  const alt = aliasOf(name.split(/\s+/)[0]);
  if (alt) {
    const byAlt = algaeByGenus[alt.toLowerCase()];
    if (byAlt) {
      const med = byAlt.filter(p => p.gs);
      const hit = pickUnit(med.length ? med : byAlt, u);
      if (hit) return { trait: hit, how: `the name ${alt}, which the trait table uses` };
    }
  }

  const g = algaeByGenus[name.toLowerCase().split(/\s+/)[0]];
  if (g) {
    const medians = g.filter(p => p.gs);
    const pool = medians.length ? medians : g;
    const hit = pickUnit(pool, u);
    if (hit) return { trait: hit,
        how: `genus median${hit.ns ? ' of ' + hit.ns + ' species' : ''}` };
  }
  return { trait: null, how: null };
}

const ALIASES = DATA.aliases || {};
function aliasOf(name) {
  const a = ALIASES[(name||'').trim().toLowerCase()];
  return a && a.toLowerCase() !== (name||'').trim().toLowerCase() ? a : null;
}

// ─── "did you mean" ───────────────────────────────────────────────────────────
// Counter sheets carry typos (Wornichinia), group names (Protozoa), life stages
// (nauplii) and stray keystrokes. Reporting "not in the table" is true but useless,
// so look for a near match and offer it.
const ALL_NAMES = (() => {
  const set = new Set();
  DATA.algae.forEach(p => { if (p.gen) set.add(p.gen); });
  DATA.algae.forEach(p => set.add(p.n.replace(/\s*\(genus median\)$/, '')));
  DATA.animals.forEach(a => { set.add(a.n); const g = a.n.split(/\s+/)[0]; if (g) set.add(g); });
  return [...set].filter(n => n && n.length > 2);
})();

// Words that are a group or a life stage rather than a taxon in the tables.
const NOT_A_TAXON = {
  'protozoa': 'a group, not a taxon: count the genus, or set the group to Other',
  'protozoan': 'a group, not a taxon',
  'ciliate': 'a group of protozoa: set the group to Other',
  'ciliates': 'a group of protozoa: set the group to Other',
  'nauplii': 'a copepod life stage: the trait tables hold adults, so set the group to Zoop and give it a body length by hand',
  'nauplius': 'a copepod life stage, see nauplii',
  'napuli': 'probably "nauplii", a copepod life stage',
  'copepodite': 'a copepod life stage, not in the trait tables',
  'copepodites': 'a copepod life stage, not in the trait tables',
  'cyclopoid': 'an order of copepods: count the genus, or set the group to Zoop and give a body length',
  'calanoid': 'an order of copepods: count the genus, or set the group to Zoop and give a body length',
  'harpacticoid': 'an order of copepods, not a taxon in the tables',
  'rotifer': 'a phylum: count the genus',
  'rotifers': 'a phylum: count the genus',
  'cladoceran': 'a group: count the genus',
  'diatom': 'a group: count the genus',
  'diatoms': 'a group: count the genus',
  'green algae': 'a group: count the genus',
  'detritus': 'not an organism',
  'unknown': 'not a taxon',
};

function editDistance(a, b) {
  const m = a.length, n = b.length;
  if (!m || !n) return Math.max(m, n);
  let prev = Array.from({length: n + 1}, (_, j) => j), cur = new Array(n + 1);
  for (let i = 1; i <= m; i++) {
    cur[0] = i;
    for (let j = 1; j <= n; j++) {
      cur[j] = Math.min(prev[j] + 1, cur[j - 1] + 1,
                        prev[j - 1] + (a[i-1] === b[j-1] ? 0 : 1));
    }
    [prev, cur] = [cur, prev];
  }
  return prev[n];
}

function didYouMean(raw) {
  const name = (raw || '').trim();
  const key = name.toLowerCase();
  if (NOT_A_TAXON[key]) return { kind: 'explain', text: NOT_A_TAXON[key] };
  if (name.length < 4 || !/[a-z]{3}/i.test(name)) return null;
  let best = null, bestD = Infinity;
  for (const cand of ALL_NAMES) {
    const c = cand.toLowerCase();
    if (Math.abs(c.length - key.length) > 4) continue;
    const d = editDistance(key, c);
    if (d < bestD) { bestD = d; best = cand; if (d === 0) break; }
  }
  if (!best) return null;
  const sim = 1 - bestD / Math.max(key.length, best.length);
  return sim >= 0.72 ? { kind: 'suggest', name: best, sim } : null;
}

function lookupAnimalLoose(rawName) {
  const a = lookupAnimal(rawName);
  if (a) return { trait: a, how: 'name' };
  const genus = (rawName||'').trim().split(/\s+/)[0].toLowerCase();
  const list = animalsByGenus[genus];
  if (list && list.length) {
    const withLen = list.filter(x => x.hi);
    const pool = withLen.length ? withLen : list;
    const lo = Math.min(...pool.map(x => x.lo || x.hi)),
          hi = Math.max(...pool.map(x => x.hi || x.lo));
    return { trait: Object.assign({}, pool[0], { lo, hi, n: rawName }),
             how: `genus, spanning ${pool.length} species` };
  }
  return { trait: null, how: null };
}

function buildLinks() {
  const nodes = [], notes = [], matched = [];
  rows.forEach(r => {
    if (!r.taxon.trim()) return;
    if (r.group === 'zoop') {
      const { trait: a, how } = lookupAnimalLoose(r.taxon);
      if (!a) notes.push({ row: r.id, taxon: r.taxon,
                           msg: 'not in the zooplankton trait table',
                           hint: didYouMean(r.taxon) });
      else if (how !== 'name') matched.push(`${r.taxon}: matched by ${how}`);
      if (a && !a.hi) notes.push({ row: r.id, taxon: r.taxon,
                                  msg: 'no body length, so it cannot be linked', hint: null });
      nodes.push({ id:r.taxon, group:'zoop', count:r.count, unit:r.unit, trait:a });
    } else if (r.group === 'phyto') {
      const { trait: p, how } = lookupAlga(r.taxon, r.unit);
      if (!p) notes.push({ row: r.id, taxon: r.taxon,
                           msg: 'not in the phytoplankton trait table',
                           hint: didYouMean(r.taxon) });
      else if (how !== 'name') matched.push(`${r.taxon}: matched by ${how}`);
      const shownUnit = p ? p.u : r.unit;
      nodes.push({ id:`${r.taxon} (${shownUnit==='filament_100um'?'filament':shownUnit})`,
                   group:'phyto', count:r.count, unit:shownUnit, trait:p });
    } else {
      nodes.push({ id:r.taxon, group:'other', count:r.count, unit:r.unit, trait:null });
      notes.push({ row:r.id, taxon:r.taxon,
                   msg:'"other" taxa have no traits yet, so they are drawn unlinked', hint:null });
    }
  });

  const links = [];
  nodes.forEach(cons => {
    const a = cons.trait;
    if (cons.group !== 'zoop' || !a) return;
    const L = a.hi || a.lo;
    const herb = (a.m === 'filter' || a.m === 'suspension' ||
                  a.t === 'herbivore' || a.t === 'omnivore');
    if (herb) {
      let upper = null, why = '';
      if (a.r) { upper = rules.rotifer_max_particle_um; why = `microphagous rotifer, up to ${upper} um`; }
      else if (L) { upper = rules.burns_slope*L + rules.burns_intercept;
                    why = `Burns 1968: 22 x ${L} mm + 4.87 = ${upper.toFixed(0)} um`; }
      if (upper !== null) {
        const lower = rules['mesh_'+(a.mesh||'default')+'_um'] ?? rules.mesh_default_um;
        const up = (a.m === 'raptorial') ? upper*rules.raptorial_particle_multiplier : upper;
        nodes.forEach(res => {
          if (res.group !== 'phyto' || !res.trait) return;
          const s = res.trait.s;
          if (s < lower || s > up) return;
          if (res.trait.f) return;                                   // a filament fouls the filter
          if (res.trait.c && s >= rules.interference_length_um) return;  // too big a colony
          links.push({ source:cons.id, target:res.id, type:'grazing',
                       why:`${why}; mesh ${lower} um` });
        });
      }
    }
    const carn = (a.m === 'raptorial' || a.t === 'carnivore' || a.t === 'omnivore');
    if (carn && L) {
      const lo = rules.prey_ratio_min*L, hi = rules.prey_ratio_max*L;
      nodes.forEach(res => {
        if (res.group !== 'zoop' || res.id === cons.id || !res.trait) return;
        const pl = res.trait.hi || res.trait.lo;
        if (pl && pl >= lo && pl <= hi) {
          links.push({ source:cons.id, target:res.id, type:'predation',
                       why:`prey ${lo.toFixed(2)}-${hi.toFixed(2)} mm of predator length` });
        }
      });
    }
  });
  return { nodes, links, notes, matched };
}

// ─── centrality ───────────────────────────────────────────────────────────────
function centrality(nodes, links) {
  const idx = {}; nodes.forEach((n,i)=>idx[n.id]=i);
  const adj = nodes.map(()=>[]);
  links.forEach(l => {
    const s = idx[l.source], t = idx[l.target];
    if (s === undefined || t === undefined) return;
    if (!adj[s].includes(t)) adj[s].push(t);
    if (!adj[t].includes(s)) adj[t].push(s);
  });
  const degree = adj.map(a=>a.length);
  // Brandes, on the undirected graph
  const bc = nodes.map(()=>0);
  for (let s=0; s<nodes.length; s++) {
    const S=[], P=nodes.map(()=>[]), sigma=nodes.map(()=>0), d=nodes.map(()=>-1);
    sigma[s]=1; d[s]=0; const Q=[s];
    while (Q.length) {
      const v=Q.shift(); S.push(v);
      adj[v].forEach(w=>{
        if (d[w]<0) { d[w]=d[v]+1; Q.push(w); }
        if (d[w]===d[v]+1) { sigma[w]+=sigma[v]; P[w].push(v); }
      });
    }
    const delta=nodes.map(()=>0);
    while (S.length) {
      const w=S.pop();
      P[w].forEach(v=>{ delta[v]+= (sigma[v]/sigma[w])*(1+delta[w]); });
      if (w!==s) bc[w]+=delta[w];
    }
  }
  const n = nodes.length;
  const norm = n>2 ? 2/((n-1)*(n-2)) : 1;     // undirected, so halve
  return nodes.map((nd,i)=>({ id:nd.id, group:nd.group,
      degree:degree[i], betweenness:+(bc[i]*norm/2).toFixed(3) }));
}

// ─── drawing ──────────────────────────────────────────────────────────────────
let sim = null;
function redraw() {
  const { nodes, links, notes, matched } = buildLinks();
  const cent = centrality(nodes, links);
  const centById = {}; cent.forEach(c=>centById[c.id]=c);

  const possible = nodes.length*(nodes.length-1);
  document.getElementById('summary').innerHTML =
    `<span class="stat"><b>${nodes.length}</b> taxa &nbsp; <b>${links.length}</b> links &nbsp; ` +
    `connectance <b>${possible? (links.length/possible).toFixed(3):'0'}</b></span>`;

  const esc = t => String(t).replace(/[&<>"]/g, ch =>
      ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[ch]));
  const noteLine = n => {
    let extra = '';
    if (n.hint && n.hint.kind === 'suggest') {
      extra = ` &mdash; did you mean <b>${esc(n.hint.name)}</b>? ` +
              `<button style="font-size:12px;padding:1px 7px;" ` +
              `onclick="applySuggestion(${n.row}, '${esc(n.hint.name).replace(/'/g, "\\'")}')">Use</button>`;
    } else if (n.hint && n.hint.kind === 'explain') {
      extra = ` &mdash; <span class="muted">${esc(n.hint.text)}</span>`;
    }
    return `<div><b>${esc(n.taxon)}</b>: ${esc(n.msg)}${extra}</div>`;
  };
  const listBox = (items, cls, heading, footer, render) => items.length ? `<div class="${cls}">` +
      `<b>${heading} (${items.length})</b>` +
      `<button onclick="copyNotes('${cls}')" style="float:right;font-size:12px;padding:2px 8px;">Copy list</button>` +
      `<div class="notelist" id="list-${cls}">` + items.map(render).join('') + `</div>` +
      (footer ? `<span class="muted">${footer}</span>` : '') + `</div>` : '';
  document.getElementById('count-warn').innerHTML =
    listBox(notes, 'warn', 'Not found or unusable', '', noteLine) +
    listBox(matched, 'from-counter', 'Matched less exactly',
            'A genus value is the median across the species in the trait table, so it is a '
            + 'rough stand-in for a species measurement.',
            m => `<div>${esc(m)}</div>`);

  document.getElementById('cent-body').innerHTML = cent
    .sort((a,b)=>b.degree-a.degree || b.betweenness-a.betweenness)
    .map(c=>`<tr><td>${c.id.replace(/</g,'&lt;')}</td><td>${c.degree}</td><td>${c.betweenness}</td></tr>`)
    .join('') || '<tr><td colspan="3" class="muted">No taxa yet.</td></tr>';

  const svg = d3.select('#graph');
  svg.selectAll('*').remove();
  if (!nodes.length) return;
  const box = document.getElementById('graph').getBoundingClientRect();
  const W = box.width, H = box.height;
  svg.attr('viewBox', `0 0 ${W} ${H}`);

  nodes.forEach(n => { n.degree = centById[n.id] ? centById[n.id].degree : 0; });
  const maxCount = d3.max(nodes, n=>n.count) || 1;
  const rOf = n => 6 + 18*Math.sqrt((n.count||0)/maxCount);
  const colour = g => g==='phyto' ? '#2e8b57' : g==='zoop' ? '#F8BD7F' : '#727171';

  const link = svg.append('g').attr('stroke','#999').attr('stroke-opacity',0.45)
    .selectAll('line').data(links).join('line')
      .attr('stroke-width', d=>d.type==='predation'?2:1)
      .attr('stroke-dasharray', d=>d.type==='predation'?'':'3,2');

  const node = svg.append('g').selectAll('g').data(nodes).join('g')
      .attr('tabindex', 0).style('cursor','pointer')
      .on('click', (e,d)=>showDetail(d, links, centById))
      .on('keydown', (e,d)=>{ if(e.key==='Enter') showDetail(d, links, centById); })
      .call(d3.drag()
        .on('start', (e,d)=>{ if(!e.active) sim.alphaTarget(0.25).restart(); d.fx=d.x; d.fy=d.y; })
        .on('drag',  (e,d)=>{ d.fx=e.x; d.fy=e.y; })
        .on('end',   (e,d)=>{ if(!e.active) sim.alphaTarget(0); d.fx=null; d.fy=null; }));
  node.append('circle').attr('r', rOf).attr('fill', d=>colour(d.group))
      .attr('stroke', d=>d.group==='zoop' ? '#b07a3c' : '#fff').attr('stroke-width',1.5)
      .append('title').text(d=>`${d.id}\ncount ${d.count}\ndegree ${centById[d.id]?.degree ?? 0}`);
  node.append('text').text(d=>d.id).attr('font-size',11).attr('dx', d=>rOf(d)+3).attr('dy',4);

  sim = d3.forceSimulation(nodes)
    .force('link', d3.forceLink(links).id(d=>d.id).distance(90).strength(0.35))
    .force('charge', d3.forceManyBody().strength(d => d.degree ? -260 : -90))
    .force('x', d3.forceX(W/2).strength(0.06))
    .force('y', d3.forceY(H/2).strength(0.06))
    .force('center', d3.forceCenter(W/2, H/2))
    .force('collide', d3.forceCollide().radius(d=>rOf(d)+14))
    .on('tick', () => {
      // keep every node inside the panel, labels included, so nothing is clipped
      nodes.forEach(d => {
        const r = rOf(d) + 4;
        const labelRoom = Math.min(8 * d.id.length + 20, W / 3);
        d.x = Math.max(r, Math.min(W - r - labelRoom, d.x));
        d.y = Math.max(r, Math.min(H - r, d.y));
      });
      link.attr('x1',d=>d.source.x).attr('y1',d=>d.source.y)
          .attr('x2',d=>d.target.x).attr('y2',d=>d.target.y);
      node.attr('transform', d=>`translate(${d.x},${d.y})`);
    });
}

function applySuggestion(rowId, name) {
  const r = rows.find(r => r.id === rowId);
  if (!r) return;
  r.taxon = name;
  const input = document.querySelector(`#r${rowId} input.taxon-in`);
  if (input) input.value = name;
  // the corrected name may belong to the other group
  const grpSel = document.querySelector(`#r${rowId} select`);
  if (lookupAnimalLoose(name).trait && !lookupAlga(name, 'cell').trait) {
    r.group = 'zoop'; if (grpSel) grpSel.value = 'zoop'; renderUnits(rowId);
  }
  redraw();
}

function copyNotes(cls) {
  const box = document.getElementById('list-' + cls);
  if (!box) return;
  const text = [...box.querySelectorAll('div')].map(d => d.textContent).join('\n');
  const done = () => { const b = box.parentElement.querySelector('button');
                       const was = b.textContent; b.textContent = 'Copied';
                       setTimeout(() => { b.textContent = was; }, 1500); };
  if (navigator.clipboard && navigator.clipboard.writeText) {
    navigator.clipboard.writeText(text).then(done, () => fallbackCopy(text, done));
  } else { fallbackCopy(text, done); }
}

function fallbackCopy(text, done) {
  const ta = document.createElement('textarea');
  ta.value = text; ta.style.position = 'fixed'; ta.style.opacity = '0';
  document.body.appendChild(ta); ta.select();
  try { document.execCommand('copy'); done(); } catch (e) { window.prompt('Copy this list:', text); }
  document.body.removeChild(ta);
}

function showDetail(d, links, centById) {
  const eats = links.filter(l=>(l.source.id||l.source)===d.id);
  const eaten = links.filter(l=>(l.target.id||l.target)===d.id);
  const name = s => s.id || s;
  document.getElementById('detail').innerHTML =
    `<b>${d.id}</b> &mdash; count ${d.count}, degree ${centById[d.id]?.degree ?? 0}, ` +
    `betweenness ${centById[d.id]?.betweenness ?? 0}<br>` +
    (eats.length ? `<b>Eats:</b> ${eats.map(l=>name(l.target)).join(', ')}<br>` : '') +
    (eaten.length ? `<b>Eaten by:</b> ${eaten.map(l=>name(l.source)).join(', ')}<br>` : '') +
    (eats[0] ? `<span class="muted">Rule: ${eats[0].why}</span>` : '');
}

// ─── loading a counter CSV ────────────────────────────────────────────────────
document.getElementById('file').addEventListener('change', ev => {
  const f = ev.target.files[0]; if (!f) return;
  const rd = new FileReader();
  rd.onload = () => { parseCsv(rd.result); };
  rd.readAsText(f);
});

function parseCsv(text) {
  const lines = text.split(/\r?\n/);
  let head = -1;
  for (let i=0;i<lines.length;i++) {
    if (/(^|,)"?taxon( name)?"?(,|$)/i.test(lines[i])) { head = i; break; }
  }
  if (head < 0) { alert('No "taxon" column found. Export from FoodWebCounter, or add taxa by hand.'); return; }
  const cols = splitCsvLine(lines[head]).map(c=>c.toLowerCase().trim());
  const ci = n => {
    let i = cols.indexOf(n);
    if (i < 0 && n === 'taxon') i = cols.indexOf('taxon name');   // older exports
    if (i < 0 && n === 'count') i = cols.findIndex(c => c.startsWith('count'));
    return i;
  };
  clearAll();
  let added = 0;
  for (let i=head+1;i<lines.length;i++) {
    if (!lines[i].trim()) continue;
    const c = splitCsvLine(lines[i]);
    const taxon = (c[ci('taxon')]||'').trim();
    if (!taxon) continue;
    const grpRaw = (c[ci('group')]||'').toLowerCase();
    // With no group column, work it out from the trait tables rather than assuming
    // phytoplankton: a bare genus like "Keratella" is a rotifer.
    const guess = () => {
      const { name } = splitUnitFromName(taxon);
      if (lookupAnimalLoose(name).trait) return 'zoop';
      if (lookupAlga(taxon, 'cell').trait) return 'phyto';
      return 'phyto';
    };
    const group = grpRaw.startsWith('zoo') ? 'zoop' : grpRaw.startsWith('phyt') ? 'phyto'
                : grpRaw ? 'other' : guess();
    let unit = (c[ci('counting_unit')]||'').trim().toLowerCase() || (group==='phyto'?'cell':'individual');
    if (unit === 'filament') unit = 'filament_100um';
    addRow(taxon, group, unit, parseFloat(c[ci('count')]) || 0);
    added++;
  }
  if (!added) alert('No taxon rows were found in that file.');
}

function splitCsvLine(line) {
  const out=[]; let cur='', q=false;
  for (let i=0;i<line.length;i++) {
    const ch=line[i];
    if (q) { if (ch==='"' && line[i+1]==='"') { cur+='"'; i++; } else if (ch==='"') q=false; else cur+=ch; }
    else if (ch==='"') q=true;
    else if (ch===',') { out.push(cur); cur=''; }
    else cur+=ch;
  }
  out.push(cur); return out;
}

// ─── a count handed over by FoodWebCounter on this same site ──────────────────
const HANDOFF_KEY = 'cfb.foodweb.count';

function loadHandoff() {
  let raw = null;
  try { raw = localStorage.getItem(HANDOFF_KEY); } catch (e) { return false; }
  if (!raw) return false;
  let p;
  try { p = JSON.parse(raw); } catch (e) { return false; }
  if (!p || !Array.isArray(p.rows) || !p.rows.length) return false;
  clearAll();
  p.rows.forEach(r => addRow(r.taxon, r.group, r.unit, r.count));
  const who = [p.meta && p.meta.name, p.meta && p.meta.location, p.meta && p.meta.date]
              .filter(Boolean).join(', ');
  document.getElementById('handoff-note').innerHTML =
    `<div class="from-counter">Loaded <b>${p.rows.length}</b> taxa from FoodWebCounter` +
    (who ? ` (${who.replace(/</g,'&lt;')})` : '') +
    `. Edit anything below; the counter is not changed.` +
    ` <button onclick="forgetHandoff()" style="margin-left:6px;">Start fresh</button></div>`;
  return true;
}

function forgetHandoff() {
  try { localStorage.removeItem(HANDOFF_KEY); } catch (e) {}
  document.getElementById('handoff-note').innerHTML = '';
  clearAll();
}

renderRules();
if (!loadHandoff()) redraw();
</script>
</body>
</html>
"""


if __name__ == "__main__":
    main()
